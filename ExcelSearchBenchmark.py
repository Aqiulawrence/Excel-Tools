"""ExcelSearch standalone: index engine, native watcher, background service and UI.

Run this file directly for the GUI. Requires Python 3.10+, openpyxl and
optionally xlrd for .xls. No other ExcelSearch Python modules are required.
The index database is stored beside this Python file.
"""
from __future__ import annotations

import contextlib
from dataclasses import dataclass, field
import hashlib
import json
import os
from pathlib import Path
import sqlite3
import statistics
import sys
import time
import uuid
import ctypes
from ctypes import wintypes
import struct
import threading
from collections import deque
import argparse
import queue

# ExcelSearchSpeedCore

sys.dont_write_bytecode = True

BASE = Path(__file__).resolve().parent

if (BASE / "_speedtest_deps").is_dir():
    sys.path.append(str(BASE / "_speedtest_deps"))

DEFAULT_DATABASE = BASE / "ExcelSearch-v3.sqlite3"

EXTENSIONS = (".xlsx", ".xlsm", ".xls")

SYSTEM_DIRS = {"$recycle.bin", "system volume information"}

INSERT_BATCH = 2000

class Cancelled(Exception):
    pass

class SourceChanged(Exception):
    pass

def check_cancel(cancel=None):
    if cancel is not None and cancel.is_set():
        raise Cancelled("操作已停止，已提交的索引保留。")

def path_key(path):
    return os.path.normcase(os.path.abspath(os.path.normpath(path)))

def database_path(database=None):
    path = Path(database or DEFAULT_DATABASE).resolve()
    if not path.is_relative_to(BASE) or path == BASE:
        raise ValueError("索引数据库必须位于 branch 文件夹内。")
    return path

def fresh_database():
    return DEFAULT_DATABASE.parent / f"index-v3-{uuid.uuid4().hex}.sqlite3"

def signature(path):
    stat = os.stat(path)
    return stat.st_size, stat.st_mtime_ns

@dataclass
class Source:
    path: str
    size: int | None = None
    mtime_ns: int | None = None
    error: str = ""

    @property
    def key(self):
        return path_key(self.path)

@dataclass
class Inventory:
    root: str
    sources: list[Source] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    inaccessible: list[str] = field(default_factory=list)

def inventory(folder, cancel=None, log=None):
    if not str(folder).strip():
        raise ValueError("请选择 Excel 文件夹。")
    root = os.path.abspath(folder)
    if not os.path.isdir(root):
        raise ValueError(f"目录不存在或无法访问：{root}")
    result = Inventory(root)
    last_progress = time.monotonic()
    directory_count = 0

    def walk_error(error):
        result.warnings.append(f"跳过目录：{error.filename}；{error}")
        if error.filename:
            result.inaccessible.append(path_key(error.filename))

    for current, directories, names in os.walk(root, onerror=walk_error, followlinks=False):
        check_cancel(cancel)
        directory_count += 1
        if log is not None and time.monotonic() - last_progress >= 0.5:
            log(f"目录扫描中：{directory_count} 个目录，已发现 {len(result.sources)} 个 Excel 文件…")
            last_progress = time.monotonic()
        # Windows protected folders are outside the useful Excel search scope.
        directories[:] = sorted(d for d in directories if d.casefold() not in SYSTEM_DIRS)
        for name in names:
            check_cancel(cancel)
            if name.startswith("~$") or not name.lower().endswith(EXTENSIONS):
                continue
            path = os.path.join(current, name)
            source = Source(path)
            try:
                source.size, source.mtime_ns = signature(path)
            except OSError as exc:
                source.error = f"{type(exc).__name__}: {exc}"
            result.sources.append(source)
    result.sources.sort(key=lambda source: source.key)
    return result

def xls_rows(sheet, workbook):
    from openpyxl.styles.numbers import is_timedelta_format
    from openpyxl.utils.datetime import MAC_EPOCH, WINDOWS_EPOCH, from_excel
    from xlrd import XL_CELL_BOOLEAN, XL_CELL_DATE, XL_CELL_ERROR, error_text_from_code
    """Normalize typed BIFF values to the same kinds of values as XLSX readers."""
    epoch = MAC_EPOCH if workbook.datemode else WINDOWS_EPOCH
    for row_number in range(sheet.nrows):
        values = sheet.row_values(row_number)
        for column, cell_type in enumerate(sheet.row_types(row_number)):
            if cell_type == XL_CELL_DATE:
                xf = workbook.xf_list[sheet.cell_xf_index(row_number, column)]
                number_format = workbook.format_map[xf.format_key].format_str
                values[column] = from_excel(
                    values[column], epoch=epoch,
                    timedelta=is_timedelta_format(number_format or ""),
                )
            elif cell_type == XL_CELL_BOOLEAN:
                values[column] = bool(values[column])
            elif cell_type == XL_CELL_ERROR:
                values[column] = error_text_from_code.get(values[column], str(values[column]))
        yield values

@contextlib.contextmanager
def workbook_sheets(path):
    # XLS support is optional; missing xlrd cannot prevent XLSX indexing.
    if path.lower().endswith(".xls"):
        try:
            from xlrd import open_workbook
        except ImportError as exc:
            raise ValueError("读取 .xls 需要 xlrd，请安装依赖。") from exc
        book = open_workbook(path, formatting_info=True)
        try:
            yield ((sheet.name, xls_rows(sheet, book)) for sheet in book.sheets())
        finally:
            book.release_resources()
        return
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("读取 .xlsx/.xlsm 需要 openpyxl，请安装依赖。") from exc
    book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        for sheet in book.worksheets:
            sheet.reset_dimensions()
        yield ((sheet.title, sheet.iter_rows(values_only=True)) for sheet in book.worksheets)
    finally:
        book.close()

SCHEMA = """
CREATE TABLE IF NOT EXISTS es3_roots (
    root_key TEXT PRIMARY KEY, path TEXT NOT NULL, updated_ns INTEGER NOT NULL DEFAULT 0
);
CREATE TABLE IF NOT EXISTS es3_files (
    id INTEGER PRIMARY KEY, root_key TEXT NOT NULL, path_key TEXT NOT NULL,
    path TEXT NOT NULL, size INTEGER, mtime_ns INTEGER, status TEXT NOT NULL,
    error TEXT NOT NULL DEFAULT '', sheets INTEGER NOT NULL DEFAULT 0,
    rows INTEGER NOT NULL DEFAULT 0, cells INTEGER NOT NULL DEFAULT 0,
    UNIQUE(root_key, path_key),
    FOREIGN KEY(root_key) REFERENCES es3_roots(root_key)
);
CREATE TABLE IF NOT EXISTS es3_cells (
    id INTEGER PRIMARY KEY, file_id INTEGER NOT NULL,
    sheet_order INTEGER NOT NULL, sheet_name TEXT NOT NULL,
    row_number INTEGER NOT NULL, column_number INTEGER NOT NULL,
    value TEXT NOT NULL, value_lower TEXT NOT NULL,
    FOREIGN KEY(file_id) REFERENCES es3_files(id) ON DELETE CASCADE
);
CREATE INDEX IF NOT EXISTS es3_cells_file ON es3_cells(file_id);
"""

FTS_SCHEMA = """
CREATE VIRTUAL TABLE IF NOT EXISTS es3_fts USING fts5(
    value_lower, content='es3_cells', content_rowid='id',
    tokenize='trigram case_sensitive 1'
);
CREATE TRIGGER IF NOT EXISTS es3_cell_insert AFTER INSERT ON es3_cells BEGIN
    INSERT INTO es3_fts(rowid, value_lower) VALUES (new.id, new.value_lower);
END;
CREATE TRIGGER IF NOT EXISTS es3_cell_delete AFTER DELETE ON es3_cells BEGIN
    INSERT INTO es3_fts(es3_fts, rowid, value_lower) VALUES ('delete', old.id, old.value_lower);
END;
"""

@dataclass
class Matches:
    count: int = 0
    digest: object = field(default_factory=hashlib.sha256)
    preview: list = field(default_factory=list)
    limit: int | None = 500

    def add(self, hit):
        self.count += 1
        encoded = json.dumps(hit, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        self.digest.update(len(encoded).to_bytes(8, "little"))
        self.digest.update(encoded)
        if self.limit is None or len(self.preview) < self.limit:
            self.preview.append(hit)

class SearchEngine:
    def __init__(self, database=None, log=print, cancel=None):
        self.database = database_path(database)
        self.log = log
        self.cancel = cancel
        self.connection = None
        self.fts = False

    def __enter__(self):
        check_cancel(self.cancel)
        self.database.parent.mkdir(parents=True, exist_ok=True)
        try:
            self.connection = sqlite3.connect(str(self.database), timeout=2)
            self.connection.row_factory = sqlite3.Row
            self.connection.execute("PRAGMA foreign_keys=ON")
            self.connection.execute("PRAGMA journal_mode=WAL")
            self.connection.execute("PRAGMA synchronous=NORMAL")
            self.connection.execute("PRAGMA cache_size=-16384")
            self.connection.executescript(SCHEMA)
            had_fts = self.connection.execute(
                "SELECT 1 FROM sqlite_master WHERE name='es3_fts'"
            ).fetchone()
            try:
                self.connection.executescript(FTS_SCHEMA)
                if not had_fts:
                    self.connection.execute("INSERT INTO es3_fts(es3_fts) VALUES ('rebuild')")
                self.fts = True
            except sqlite3.OperationalError as exc:
                if "no such module" not in str(exc) and "tokenizer" not in str(exc):
                    raise
                self.log("当前 SQLite 不支持 trigram，使用 SQLite 子串查询，测速会标注。")
            self.connection.commit()
            return self
        except Exception:
            self.__exit__(None, None, None)
            raise

    def __exit__(self, *args):
        if self.connection is not None:
            self.connection.close()
            self.connection = None

    def _file_record(self, root_key, source, status, error=""):
        self.connection.execute(
            "INSERT INTO es3_files(root_key,path_key,path,size,mtime_ns,status,error) "
            "VALUES (?,?,?,?,?,?,?) ON CONFLICT(root_key,path_key) DO UPDATE SET "
            "path=excluded.path,size=excluded.size,mtime_ns=excluded.mtime_ns,"
            "status=excluded.status,error=excluded.error,sheets=0,rows=0,cells=0",
            (root_key, source.key, source.path, source.size, source.mtime_ns, status, error),
        )
        return self.connection.execute(
            "SELECT id FROM es3_files WHERE root_key=? AND path_key=?", (root_key, source.key)
        ).fetchone()[0]

    def _index_file(self, root_key, source):
        check_cancel(self.cancel)
        if source.error:
            raise OSError(source.error)
        with self.connection:
            file_id = self._file_record(root_key, source, "ok")
            self.connection.execute("DELETE FROM es3_cells WHERE file_id=?", (file_id,))
            sheet_count = row_count = cell_count = 0
            batch = []
            insert_sql = ("INSERT INTO es3_cells(file_id,sheet_order,sheet_name,row_number,"
                          "column_number,value,value_lower) VALUES (?,?,?,?,?,?,?)")
            with workbook_sheets(source.path) as sheets:
                for sheet_order, (name, rows) in enumerate(sheets):
                    sheet_count += 1
                    for row_number, row in enumerate(rows, 1):
                        check_cancel(self.cancel)
                        row_count += 1
                        for column, value in enumerate(row, 1):
                            if value is None or value == "":
                                continue
                            value = str(value)
                            batch.append((file_id, sheet_order, name, row_number, column, value, value.lower()))
                            cell_count += 1
                            if len(batch) >= INSERT_BATCH:
                                check_cancel(self.cancel)
                                self.connection.executemany(insert_sql, batch)
                                batch.clear()
                if batch:
                    self.connection.executemany(insert_sql, batch)
            if signature(source.path) != (source.size, source.mtime_ns):
                raise SourceChanged("文件在读取期间发生变化，将在下次更新时重试。")
            self.connection.execute("UPDATE es3_files SET sheets=?,rows=?,cells=? WHERE id=?",
                                    (sheet_count, row_count, cell_count, file_id))

    def prepare(self, folder, retry_errors=False):
        """One inventory, per-file transactions, no global all-files-ready gate."""
        started = time.perf_counter()
        self.log("扫描文件夹并检查增量变化…")
        snapshot = inventory(folder, self.cancel, self.log)
        root_key = path_key(snapshot.root)
        for warning in snapshot.warnings[:30]:
            self.log(warning)
        if len(snapshot.warnings) > 30:
            self.log(f"另有 {len(snapshot.warnings) - 30} 个目录访问失败。")
        with self.connection:
            self.connection.execute("INSERT INTO es3_roots(root_key,path) VALUES (?,?) "
                                    "ON CONFLICT(root_key) DO UPDATE SET path=excluded.path",
                                    (root_key, snapshot.root))
        previous = {row["path_key"]: row for row in self.connection.execute(
            "SELECT * FROM es3_files WHERE root_key=?", (root_key,))}
        updated = reused = 0
        self.log(f"发现 {len(snapshot.sources)} 个 Excel 文件。")
        last_progress = time.monotonic()
        for position, source in enumerate(snapshot.sources, 1):
            check_cancel(self.cancel)
            old = previous.get(source.key)
            same = old is not None and (old["size"], old["mtime_ns"]) == (source.size, source.mtime_ns)
            if (same and not source.error and old["status"] != "inaccessible"
                    and (old["status"] == "ok" or not retry_errors)):
                reused += 1
            else:
                try:
                    self._index_file(root_key, source)
                    updated += 1
                except Cancelled:
                    raise
                except sqlite3.Error:
                    # Database failures must not be misclassified as bad Excel files.
                    raise
                except Exception as exc:
                    error = f"{type(exc).__name__}: {exc}"
                    with self.connection:
                        file_id = self._file_record(root_key, source, "error", error)
                        self.connection.execute("DELETE FROM es3_cells WHERE file_id=?", (file_id,))
                    self.log(f"跳过文件：{source.path}；{error}")
            if time.monotonic() - last_progress >= 0.5 or position == len(snapshot.sources):
                self.log(f"处理进度：{position}/{len(snapshot.sources)}，已用 {time.perf_counter() - started:.1f} 秒")
                last_progress = time.monotonic()
        current_keys = {source.key for source in snapshot.sources}
        inaccessible = {
            key for key in previous
            if any(key == folder or key.startswith(os.path.join(folder, ""))
                   for folder in snapshot.inaccessible)
        }
        removed = set(previous) - current_keys - inaccessible
        with self.connection:
            for key in inaccessible:
                check_cancel(self.cancel)
                self.connection.execute(
                    "UPDATE es3_files SET status='inaccessible',error=? WHERE root_key=? AND path_key=?",
                    ("所在目录暂时无法访问；保留缓存，恢复后重新核对。", root_key, key),
                )
            for key in removed:
                check_cancel(self.cancel)
                self.connection.execute("DELETE FROM es3_files WHERE root_key=? AND path_key=?", (root_key, key))
            self.connection.execute("UPDATE es3_roots SET updated_ns=? WHERE root_key=?", (time.time_ns(), root_key))
        counts = dict(self.connection.execute(
            "SELECT count(*) AS files,coalesce(sum(sheets),0) AS sheets,"
            "coalesce(sum(rows),0) AS rows,coalesce(sum(cells),0) AS cells "
            "FROM es3_files WHERE root_key=? AND status='ok'", (root_key,)).fetchone())
        failed = self.connection.execute("SELECT count(*) FROM es3_files WHERE root_key=? AND status!='ok'",
                                         (root_key,)).fetchone()[0]
        elapsed = time.perf_counter() - started
        self.log(f"索引可用：成功 {counts['files']} 个，跳过 {failed} 个；更新 {updated} 个，"
                 f"复用 {reused} 个，移除 {len(removed)} 个；准备耗时 {elapsed:.3f} 秒。")
        if not counts["files"]:
            self.log("当前范围没有可搜索的 Excel 数据；可检查上面的目录或文件错误。")
        self.log(f"索引：{self.database}")
        size = sum(path.stat().st_size for path in (self.database, Path(str(self.database) + "-wal")) if path.exists())
        return {"seconds": elapsed, "ready": True, "counts": counts, "skipped": failed,
                "updated": updated, "reused": reused, "removed": len(removed),
                "directory_errors": len(snapshot.warnings), "database_mib": size / 1024**2}

    def search(self, folder, term, limit=500):
        lowered = term.lower()
        use_fts = self.fts and len(lowered) >= 3 and "\x00" not in lowered
        select = ("SELECT f.path,c.sheet_name,c.row_number,c.column_number,c.value "
                  "FROM es3_cells c JOIN es3_files f ON f.id=c.file_id ")
        if use_fts:
            select += "JOIN es3_fts ON es3_fts.rowid=c.id "
            where = "es3_fts MATCH ? AND "
            parameters = ['"' + lowered.replace('"', '""') + '"']
        else:
            where, parameters = "", []
        select += ("WHERE " + where + "f.root_key=? AND f.status='ok' AND instr(c.value_lower,?)>0 "
                   "ORDER BY f.path_key,c.sheet_order,c.row_number,c.column_number")
        parameters += [path_key(folder), lowered]
        matches = Matches(limit=limit)
        self.connection.set_progress_handler(lambda: int(self.cancel is not None and self.cancel.is_set()), 2000)
        try:
            for row in self.connection.execute(select, parameters):
                check_cancel(self.cancel)
                matches.add(tuple(row))
        except sqlite3.OperationalError:
            check_cancel(self.cancel)
            raise
        finally:
            self.connection.set_progress_handler(None, 0)
        return matches, "FTS5 trigram" if use_fts else "SQLite 子串扫描"

    def scan_snapshot(self, folder, term):
        matches = Matches(limit=0)
        errors = []
        files = list(self.connection.execute("SELECT * FROM es3_files WHERE root_key=? AND status='ok' ORDER BY path_key",
                                              (path_key(folder),)))
        for source in files:
            check_cancel(self.cancel)
            try:
                expected = source["size"], source["mtime_ns"]
                if signature(source["path"]) != expected:
                    raise SourceChanged("源文件已变化")
                with workbook_sheets(source["path"]) as sheets:
                    for name, rows in sheets:
                        for row_number, row in enumerate(rows, 1):
                            check_cancel(self.cancel)
                            for column, value in enumerate(row, 1):
                                if value is not None and value != "" and term.lower() in str(value).lower():
                                    matches.add((source["path"], name, row_number, column, str(value)))
                if signature(source["path"]) != expected:
                    raise SourceChanged("源文件在扫描期间变化")
            except Cancelled:
                raise
            except Exception as exc:
                errors.append(f"{source['path']}；{type(exc).__name__}: {exc}")
        return matches, errors

    def benchmark(self, folder, term, repeat=1, compare=False, limit=500, prepared=None):
        if not term.strip():
            raise ValueError("请输入非空搜索关键词。")
        if not isinstance(repeat, int) or not 1 <= repeat <= 1000:
            raise ValueError("重复次数必须为 1～1000 的整数。")
        if prepared is None:
            prepared = self.prepare(folder)
        timings = []
        first_digest = None
        # One stable database snapshot for all repetitions and comparison scope.
        self.connection.execute("BEGIN")
        try:
            for run in range(repeat):
                check_cancel(self.cancel)
                started = time.perf_counter()
                matches, backend = self.search(folder, term, limit=limit)
                elapsed = time.perf_counter() - started
                timings.append(elapsed)
                digest = matches.digest.digest()
                if first_digest is not None and digest != first_digest:
                    raise RuntimeError("同一索引快照的查询结果不一致。")
                first_digest = digest
                self.log(f"索引查询 {run + 1}/{repeat}：{elapsed * 1000:.3f} ms，{matches.count} 个匹配单元格")
            median = statistics.median(timings)
            report = {"keyword": term, "matches": matches.count, "repeat": repeat,
                      "backend": backend, "prepare_seconds": prepared["seconds"],
                      "files": prepared["counts"]["files"], "skipped_files": prepared["skipped"],
                      "index_seconds": timings, "first_ms": timings[0] * 1000,
                      "median_ms": median * 1000, "min_ms": min(timings) * 1000,
                      "mean_ms": statistics.mean(timings) * 1000,
                      "comparison": "not_requested"}
            self.log(f"{backend}：首次 {report['first_ms']:.3f} ms，中位数 {report['median_ms']:.3f} ms；"
                     "目录扫描与增量更新计入准备耗时，不计入查询耗时。")
            if compare:
                self.log("正在原始扫描相同的成功索引文件，核对全部匹配…")
                started = time.perf_counter()
                scanned, errors = self.scan_snapshot(folder, term)
                elapsed = time.perf_counter() - started
                report["scan_seconds"] = elapsed
                if errors:
                    report["comparison"] = "source_changed_or_unreadable"
                    report["comparison_errors"] = errors
                    self.log("源文件已变化或无法读取，本次不计算加速比；索引快照搜索结果保留。")
                    for error in errors[:20]:
                        self.log(error)
                elif scanned.count != matches.count or scanned.digest.digest() != first_digest:
                    report["comparison"] = "mismatch"
                    self.log("结果校验未通过，本次不计算加速比。请更新索引后重试。")
                else:
                    report.update(comparison="equal", speedup=elapsed / max(median, 1e-12))
                    self.log(f"原始扫描 {elapsed * 1000:.3f} ms；全部结果一致；加速比 {report['speedup']:.2f} 倍")
            return report, matches.preview
        finally:
            self.connection.rollback()

def build_index(folder, database=None, log=print, cancel=None):
    with SearchEngine(database, log, cancel) as engine:
        return engine.prepare(folder, retry_errors=True)

def benchmark(folder, database=None, term="", repeat=1, compare=False, log=print, cancel=None, limit=500):
    # Validate before opening or building an index.
    if not term.strip():
        raise ValueError("请输入非空搜索关键词。")
    if not isinstance(repeat, int) or not 1 <= repeat <= 1000:
        raise ValueError("重复次数必须为 1～1000 的整数。")
    with SearchEngine(database, log, cancel) as engine:
        return engine.benchmark(folder, term, repeat, compare, limit)

def indexed_search(folder, database, term):
    with SearchEngine(database, log=lambda text: None) as engine:
        engine.prepare(folder)
        return engine.search(folder, term, limit=None)[0].preview

# ExcelSearchWatcher

class DirectoryWatcher:
    def __init__(self, folder, changed, invalidated):
        self.folder = os.path.abspath(folder)
        self.changed, self.invalidated = changed, invalidated
        self.stopping = threading.Event()
        self.ready = threading.Event()
        self.condition = threading.Condition()
        self.requested = self.completed = 0
        self.error = None
        self.thread = None

    def start(self):
        if os.name != "nt":
            raise RuntimeError("实时监听需要 Windows。")
        self.thread = threading.Thread(target=self._run, name="ExcelDirectoryWatcher", daemon=False)
        self.thread.start()
        if not self.ready.wait(5):
            self.stop()
            raise RuntimeError("监听启动超时。")
        if self.error:
            self.stop()
            raise RuntimeError(f"监听启动失败：{self.error}")

    def flush(self, cancel=None):
        with self.condition:
            target = self.requested = self.requested + 1
            deadline = time.monotonic() + 5
            while self.completed < target:
                check_cancel(cancel)
                if self.error or self.stopping.is_set():
                    raise RuntimeError(f"监听已中断：{self.error or '已停止'}")
                if time.monotonic() > deadline:
                    raise RuntimeError("等待监听同步超时。")
                self.condition.wait(0.02)

    def stop(self):
        self.stopping.set()
        if self.thread and self.thread is not threading.current_thread():
            self.thread.join(5)
            if self.thread.is_alive():
                raise RuntimeError("目录监听未能退出。")

    def _run(self):
        api = None
        handle = event = None
        armed = False
        overlapped = None
        try:
            api = ctypes.WinDLL("kernel32", use_last_error=True)

            class OVERLAPPED(ctypes.Structure):
                _fields_ = [("Internal", ctypes.c_size_t), ("InternalHigh", ctypes.c_size_t),
                            ("Offset", wintypes.DWORD), ("OffsetHigh", wintypes.DWORD),
                            ("hEvent", wintypes.HANDLE)]

            api.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD,
                                       ctypes.c_void_p, wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
            api.CreateFileW.restype = wintypes.HANDLE
            api.CreateEventW.argtypes = [ctypes.c_void_p, wintypes.BOOL, wintypes.BOOL, wintypes.LPCWSTR]
            api.CreateEventW.restype = wintypes.HANDLE
            api.ReadDirectoryChangesW.argtypes = [wintypes.HANDLE, ctypes.c_void_p, wintypes.DWORD,
                wintypes.BOOL, wintypes.DWORD, ctypes.c_void_p, ctypes.POINTER(OVERLAPPED), ctypes.c_void_p]
            api.ReadDirectoryChangesW.restype = wintypes.BOOL
            api.GetOverlappedResult.argtypes = [wintypes.HANDLE, ctypes.POINTER(OVERLAPPED),
                                               ctypes.POINTER(wintypes.DWORD), wintypes.BOOL]
            api.GetOverlappedResult.restype = wintypes.BOOL
            api.WaitForSingleObject.argtypes = [wintypes.HANDLE, wintypes.DWORD]
            api.WaitForSingleObject.restype = wintypes.DWORD
            api.ResetEvent.argtypes = [wintypes.HANDLE]
            api.CancelIoEx.argtypes = [wintypes.HANDLE, ctypes.POINTER(OVERLAPPED)]
            api.CloseHandle.argtypes = [wintypes.HANDLE]
            handle = api.CreateFileW(self.folder, 1, 7, None, 3, 0x42000000, None)
            if handle == ctypes.c_void_p(-1).value:
                handle = None
                raise ctypes.WinError(ctypes.get_last_error())
            event = api.CreateEventW(None, True, False, None)
            if not event:
                raise ctypes.WinError(ctypes.get_last_error())
            buffer = ctypes.create_string_buffer(64 * 1024)
            overlapped = OVERLAPPED()
            overlapped.hEvent = event

            def arm():
                api.ResetEvent(event)
                if not api.ReadDirectoryChangesW(handle, buffer, len(buffer), True, 0x5B,
                                                 None, ctypes.byref(overlapped), None):
                    if ctypes.get_last_error() != 997:
                        raise ctypes.WinError(ctypes.get_last_error())

            arm()
            armed = True
            self.ready.set()
            while not self.stopping.is_set():
                result = api.WaitForSingleObject(event, 10)
                if result == 0:
                    count = wintypes.DWORD()
                    ok = api.GetOverlappedResult(handle, ctypes.byref(overlapped), ctypes.byref(count), False)
                    armed = False
                    if not ok and ctypes.get_last_error() != 1022:
                        raise ctypes.WinError(ctypes.get_last_error())
                    data = buffer.raw[:count.value] if ok else b""
                    arm()  # Re-arm before parsing/callbacks; OS buffers new changes.
                    armed = True
                    if not data:
                        self.invalidated("监听通知溢出，后台补扫。")
                    offset = 0
                    while offset + 12 <= len(data):
                        next_offset, action, length = struct.unpack_from("<III", data, offset)
                        if offset + 12 + length > len(data):
                            self.invalidated("监听通知不完整，后台补扫。")
                            break
                        name = data[offset + 12:offset + 12 + length].decode("utf-16-le")
                        self.changed(os.path.join(self.folder, name))
                        if not next_offset:
                            break
                        if next_offset < 12:
                            raise RuntimeError("无效目录通知。")
                        offset += next_offset
                elif result != 258:
                    raise ctypes.WinError(ctypes.get_last_error())
                if api.WaitForSingleObject(event, 0) == 258:
                    with self.condition:
                        self.completed = self.requested
                        self.condition.notify_all()
        except Exception as exc:
            self.error = str(exc)
            if not self.stopping.is_set():
                self.invalidated(f"监听中断：{exc}；暂停搜索，重连后补扫。")
        finally:
            self.ready.set()
            if api and handle:
                if armed:
                    api.CancelIoEx(handle, ctypes.byref(overlapped))
                    count = wintypes.DWORD()
                    api.GetOverlappedResult(handle, ctypes.byref(overlapped), ctypes.byref(count), True)
                api.CloseHandle(handle)
            if api and event:
                api.CloseHandle(event)
            with self.condition:
                self.condition.notify_all()

# ExcelSearchLiveIndex

class StopToken(threading.Event):
    def __init__(self, shutdown):
        super().__init__()
        self.shutdown = shutdown

    def is_set(self):
        return self.shutdown.is_set() or super().is_set()

class LiveIndexService:
    def __init__(self, folder, database, emit, watcher_factory=DirectoryWatcher):
        self.folder = os.path.abspath(folder)
        self.database = database_path(database)
        self.emit, self.watcher_factory = emit, watcher_factory
        self.condition = threading.Condition()
        self.shutdown = threading.Event()
        self.operation_cancel = StopToken(self.shutdown)
        self.pending = set()
        self.rescan = True
        self.requests = deque()
        self.retries = {}
        self.watcher = self.engine = None
        self.last_report = self.last_error = None
        self.thread = threading.Thread(target=self._run, name="ExcelLiveIndex", daemon=False)

    def start(self):
        self.thread.start()

    def log(self, text):
        self.emit("log", text)

    def state(self, text):
        self.emit("live", text)

    def changed(self, path):
        parts = Path(os.path.relpath(path, self.folder)).parts
        if any(part.casefold() in SYSTEM_DIRS for part in parts) or os.path.basename(path).startswith("~$"):
            return
        if os.path.isfile(path) and not path.lower().endswith(EXTENSIONS):
            return
        with self.condition:
            self.pending.add(path_key(path))
            if len(self.pending) > 10000:
                self.pending.clear()
                self.rescan = True
            self.condition.notify_all()

    def invalidated(self, reason):
        self.log(reason)
        self.state("监听待恢复 / 正在补扫；搜索将等待同步。")
        with self.condition:
            self.rescan = True
            self.condition.notify_all()

    def request(self, action, **options):
        with self.condition:
            self.requests.append((action, options))
            self.condition.notify_all()

    def _emit_done(self, options, outcome):
        # Optional request identity lets clients ignore a cancelled request's
        # late completion while another search is already queued.
        if "search_cancel" in options:
            outcome = {"search_cancel": options["search_cancel"], "outcome": outcome}
        self.emit("done", outcome)

    def cancel(self):
        self.operation_cancel.set()
        with self.condition:
            queued = bool(self.requests)
            self.requests.clear()
            self.condition.notify_all()
        if queued:
            self.emit("done", "操作已停止；下次搜索前会继续同步。")

    def stop(self):
        self.shutdown.set()
        self.cancel()

    def join(self, timeout=None):
        self.thread.join(timeout)

    def _connect_watcher(self):
        if not os.path.isdir(self.folder):
            raise ValueError(f"目录不存在或无法访问：{self.folder}")
        if self.watcher is not None and not self.watcher.error and not self.watcher.stopping.is_set():
            return
        if self.watcher:
            self.watcher.stop()
        self.watcher = self.watcher_factory(self.folder, self.changed, self.invalidated)
        self.watcher.start()
        self.rescan = True
        self.log("目录监听已开启（无需管理员权限）；开始后台核对索引。")

    def _counts(self, elapsed=0):
        key = path_key(self.folder)
        counts = dict(self.engine.connection.execute(
            "SELECT count(*) AS files,coalesce(sum(sheets),0) AS sheets,"
            "coalesce(sum(rows),0) AS rows,coalesce(sum(cells),0) AS cells "
            "FROM es3_files WHERE root_key=? AND status='ok'", (key,)).fetchone())
        failed = self.engine.connection.execute(
            "SELECT count(*) FROM es3_files WHERE root_key=? AND status!='ok'", (key,)).fetchone()[0]
        self.last_report = {"seconds": elapsed, "counts": counts, "skipped": failed, "ready": True}
        return self.last_report

    def _set_retry(self, path, attempt=0):
        path = path_key(path)
        if attempt < 3:
            self.retries[path] = (time.monotonic() + (1, 3, 10)[attempt], attempt + 1)
        else:
            self.retries.pop(path, None)

    def _update_file(self, path, attempt=0):
        check_cancel(self.operation_cancel)
        path = path_key(path)
        key = path_key(self.folder)
        connection = self.engine.connection
        source = Source(path)
        try:
            source.size, source.mtime_ns = signature(path)
        except FileNotFoundError:
            with connection:
                connection.execute("DELETE FROM es3_files WHERE root_key=? AND path_key=?", (key, path_key(path)))
            self.retries.pop(path, None)
            return
        except OSError as exc:
            source.error = str(exc)
        try:
            self.engine._index_file(key, source)
            self.retries.pop(path, None)
        except (Cancelled, sqlite3.Error):
            raise
        except Exception as exc:
            with connection:
                file_id = self.engine._file_record(key, source, "error", f"{type(exc).__name__}: {exc}")
                connection.execute("DELETE FROM es3_cells WHERE file_id=?", (file_id,))
            self.log(f"文件暂不可读：{path}；{exc}")
            self._set_retry(path, attempt)

    def _apply_paths(self, paths):
        if not paths:
            return
        previous = [row[0] for row in self.engine.connection.execute(
            "SELECT path FROM es3_files WHERE root_key=?", (path_key(self.folder),))]
        files = set()
        inaccessible = set()
        for path in paths:
            check_cancel(self.operation_cancel)
            prefix = path_key(path) + os.sep
            descendants = [old for old in previous if path_key(old).startswith(prefix)]
            if os.path.isdir(path) or descendants:
                files.update(descendants)
                if os.path.isdir(path):
                    snapshot = inventory(path, self.operation_cancel)
                    files.update(source.path for source in snapshot.sources)
                    inaccessible.update(snapshot.inaccessible)
                    for warning in snapshot.warnings[:10]:
                        self.log(warning)
            elif path.lower().endswith(EXTENSIONS):
                files.add(path)
        for path in sorted(files):
            normalized = path_key(path)
            if any(normalized == folder or normalized.startswith(os.path.join(folder, ""))
                   for folder in inaccessible):
                check_cancel(self.operation_cancel)
                with self.engine.connection:
                    self.engine.connection.execute(
                        "UPDATE es3_files SET status='inaccessible',error=? WHERE root_key=? AND path_key=?",
                        ("所在目录暂时无法访问；保留缓存，恢复后重新核对。", path_key(self.folder), normalized),
                    )
                self.retries.pop(normalized, None)
            else:
                self._update_file(normalized)
        if files:
            self.log(f"后台增量更新：处理 {len(files)} 个 Excel 路径。")

    def _synchronize(self, force=False):
        self._connect_watcher()
        if force:
            self.rescan = True
        self.state("正在更新索引，搜索将等待本轮同步完成…")
        started = time.perf_counter()
        deadline = time.monotonic() + 30
        while True:
            check_cancel(self.operation_cancel)
            self.watcher.flush(self.operation_cancel)
            with self.condition:
                full_scan, paths = self.rescan, self.pending
                self.rescan, self.pending = False, set()
            try:
                if full_scan:
                    self.last_report = self.engine.prepare(self.folder, retry_errors=True)
                    inaccessible = self.engine.connection.execute(
                        "SELECT path FROM es3_files WHERE root_key=? AND status='inaccessible'",
                        (path_key(self.folder),),
                    )
                    for row in inaccessible:
                        self.retries.pop(path_key(row[0]), None)
                    failed = self.engine.connection.execute(
                        "SELECT path FROM es3_files WHERE root_key=? AND status='error'", (path_key(self.folder),))
                    for row in failed:
                        if path_key(row[0]) not in self.retries:
                            self._set_retry(row[0])
                else:
                    self._apply_paths(paths)
                self.watcher.flush(self.operation_cancel)
                with self.condition:
                    has_changes = self.rescan or bool(self.pending)
                if not has_changes:
                    break
                if time.monotonic() >= deadline:
                    raise RuntimeError("文件持续变化，暂未完成同步；请等待后台更新后再搜索。")
            except Exception:
                with self.condition:
                    self.pending.update(paths)
                    self.rescan = self.rescan or full_scan
                raise
        self.last_error = None
        report = self._counts(time.perf_counter() - started)
        suffix = f"；{report['skipped']} 个文件不可读，未纳入搜索" if report["skipped"] else ""
        self.state(f"持续监听中 · 已处理当前收到的变化 · {time.strftime('%H:%M:%S')}"
                   f" · {report['counts']['files']} 个成功文件{suffix}")
        return report

    def _execute(self, action, options):
        started = time.perf_counter()
        prepared = self._synchronize(force=action == "build")
        prepared["seconds"] = time.perf_counter() - started
        if action == "build":
            return f"索引更新完成：成功 {prepared['counts']['files']} 个，跳过 {prepared['skipped']} 个。"
        self.log(f"搜索前监听同步等待 {prepared['seconds'] * 1000:.3f} ms；日常查询不遍历整盘。")
        report, hits = self.engine.benchmark(self.folder, options["term"], options["repeat"],
                                             options["compare"], prepared=prepared)
        self.emit("results", (report, hits))
        with self.condition:
            changing = bool(self.pending) or self.rescan
        if changing:
            self.log("查询期间又收到文件变化；本次结果对应同步完成时的索引快照，后台将继续更新。")
        text = f"搜索完成：{report['matches']} 个匹配，查询中位数 {report['median_ms']:.3f} ms。"
        if report["comparison"] in ("mismatch", "source_changed_or_unreadable"):
            text += " 原始扫描对比无效，详情见日志。"
        return text

    def _run(self):
        try:
            while not self.shutdown.is_set():
                action = None
                try:
                    self.operation_cancel.clear()
                    with self.condition:
                        if self.requests:
                            action, options = self.requests.popleft()
                    self._connect_watcher()
                    if self.engine is None:
                        self.engine = SearchEngine(self.database, self.log, self.operation_cancel)
                        self.engine.__enter__()
                    with self.condition:
                        needs_sync = self.rescan or bool(self.pending)
                    if action:
                        self._emit_done(options, self._execute(action, options))
                    elif needs_sync:
                        if self.shutdown.wait(0.15):
                            break
                        self._synchronize()
                    else:
                        due = [(path, attempt) for path, (when, attempt) in self.retries.items()
                               if when <= time.monotonic()]
                        for path, attempt in due:
                            self._update_file(path, attempt)
                        if due:
                            self._synchronize()
                        with self.condition:
                            if not self.requests and not self.pending and not self.rescan:
                                self.condition.wait(0.1)
                except Cancelled as exc:
                    self.state("更新已停止；下次搜索前会继续同步，当前不标记为最新。")
                    self.rescan = True
                    if action:
                        self._emit_done(options, str(exc))
                    with self.condition:
                        while not self.shutdown.is_set() and not self.requests:
                            self.condition.wait(0.1)
                except Exception as exc:
                    self.last_error = str(exc)
                    self.state(f"索引同步未完成：{exc}；将自动重试。")
                    self.log(f"后台更新失败：{type(exc).__name__}: {exc}")
                    self.rescan = True
                    if self.engine:
                        self.engine.__exit__(None, None, None)
                        self.engine = None
                    if action:
                        self._emit_done(options, f"搜索/更新未完成：{exc}")
                    deadline = time.monotonic() + 3
                    with self.condition:
                        while not self.shutdown.is_set() and not self.requests and time.monotonic() < deadline:
                            self.condition.wait(0.1)
        finally:
            if self.engine:
                self.engine.__exit__(None, None, None)
            if self.watcher:
                self.watcher.stop()
            self.emit("stopped", None)

# ExcelSearchSpeedTest

sys.dont_write_bytecode = True

def friendly_error(exc):
    if isinstance(exc, sqlite3.DatabaseError):
        detail = str(exc)
        if "locked" in detail or "busy" in detail:
            return "索引被另一个任务占用，请等待其结束后重试。"
        if "malformed" in detail or "not a database" in detail:
            return "索引数据库已损坏。点击“新建索引”即可重新建立，原数据库保留。"
        return f"数据库操作失败：{detail}。可点击“新建索引”重试。"
    if isinstance(exc, PermissionError):
        return f"没有读写权限：{exc}"
    return f"{type(exc).__name__}: {exc}"

class SpeedTestWindow:
    """All Tk access stays on the UI thread; one cancellable worker at a time."""

    def __init__(self, root, folder="D:/", database=None):
        import tkinter as tk
        from tkinter import ttk

        self.root = root
        self.tk = tk
        self.ttk = ttk
        self.events = queue.Queue()
        self.service = None
        self.generation = 0
        self.restart_id = None
        self.queued_request = None
        self.worker = None
        self.busy = False
        self.closing = False
        self.last_report = None
        self.folder = tk.StringVar(value=folder)
        self.database = tk.StringVar(value=str(database_path(database)))
        self.term = tk.StringVar()
        self.repeat = tk.StringVar(value="1")
        self.compare = tk.BooleanVar(value=False)
        self.status = tk.StringVar(value="后台正在更新索引，可输入关键词；搜索会等待同步完成。")
        self.live_status = tk.StringVar(value="启动扫描未完成；搜索将等待同步。")
        self.progress = tk.StringVar(value="使用固定数据库，重启后继续复用；默认目录 D:/。")
        root.title("ExcelSearch 索引速度测试 v4 · 实时监听")
        root.geometry("1120x800")
        root.minsize(850, 580)
        frame = ttk.Frame(root, padding=10)
        frame.pack(fill="both", expand=True)
        frame.columnconfigure(1, weight=1)
        frame.rowconfigure(6, weight=1)
        ttk.Label(frame, text="Excel 文件夹").grid(row=0, column=0, sticky="w", padx=4)
        folder_entry = ttk.Entry(frame, textvariable=self.folder)
        folder_entry.grid(row=0, column=1, sticky="ew")
        choose = ttk.Button(frame, text="选择目录", command=self.choose_folder)
        choose.grid(row=0, column=2, padx=4)
        ttk.Label(frame, text="索引数据库").grid(row=1, column=0, sticky="w", padx=4)
        ttk.Label(frame, textvariable=self.database, wraplength=900).grid(
            row=1, column=1, columnspan=2, sticky="w", pady=5)
        ttk.Label(frame, text="关键词").grid(row=2, column=0, sticky="w", padx=4)
        term_entry = ttk.Entry(frame, textvariable=self.term)
        term_entry.grid(row=2, column=1, sticky="ew")
        term_entry.bind("<Return>", lambda event: self.start("search"))

        controls = ttk.Frame(frame)
        controls.grid(row=3, column=0, columnspan=3, sticky="ew", pady=8)
        ttk.Label(controls, text="重复次数").pack(side="left")
        repeat_entry = ttk.Spinbox(controls, from_=1, to=1000, textvariable=self.repeat, width=6)
        repeat_entry.pack(side="left", padx=4)
        comparison = ttk.Checkbutton(controls, text="原始扫描对比", variable=self.compare)
        comparison.pack(side="left", padx=6)
        build = ttk.Button(controls, text="建立/更新索引", command=lambda: self.start("build"))
        build.pack(side="left", padx=4)
        search = ttk.Button(controls, text="搜索测速", command=lambda: self.start("search"))
        search.pack(side="left", padx=4)
        new = ttk.Button(controls, text="新建索引", command=self.new_database)
        new.pack(side="left", padx=4)
        self.stop_button = ttk.Button(controls, text="停止", command=self.stop, state="disabled")
        self.stop_button.pack(side="left", padx=4)
        self.inputs = [folder_entry, choose, term_entry, repeat_entry, comparison, build, search, new]
        ttk.Label(frame, textvariable=self.status, wraplength=1050).grid(
            row=4, column=0, columnspan=3, sticky="w")
        ttk.Label(frame, textvariable=self.live_status, wraplength=1050).grid(
            row=5, column=0, columnspan=3, sticky="w", pady=4)

        pane = ttk.Panedwindow(frame, orient="vertical")
        pane.grid(row=6, column=0, columnspan=3, sticky="nsew")
        log_frame = ttk.Frame(pane)
        log_frame.rowconfigure(0, weight=1)
        log_frame.columnconfigure(0, weight=1)
        self.output = tk.Text(log_frame, height=13, wrap="word", state="disabled")
        self.output.grid(row=0, column=0, sticky="nsew")
        log_scroll = ttk.Scrollbar(log_frame, orient="vertical", command=self.output.yview)
        log_scroll.grid(row=0, column=1, sticky="ns")
        self.output.configure(yscrollcommand=log_scroll.set)
        pane.add(log_frame, weight=1)

        results_frame = ttk.Frame(pane)
        results_frame.rowconfigure(0, weight=1)
        results_frame.columnconfigure(0, weight=1)
        self.table = ttk.Treeview(results_frame, columns=("file", "sheet", "row", "column", "value"),
                                 show="headings")
        for column, title, width in zip(("file", "sheet", "row", "column", "value"),
                                        ("文件", "工作表", "行", "列", "单元格内容"),
                                        (370, 120, 65, 65, 380)):
            self.table.heading(column, text=title)
            self.table.column(column, width=width, minwidth=50, stretch=False)
        self.table.grid(row=0, column=0, sticky="nsew")
        vertical = ttk.Scrollbar(results_frame, orient="vertical", command=self.table.yview)
        vertical.grid(row=0, column=1, sticky="ns")
        horizontal = ttk.Scrollbar(results_frame, orient="horizontal", command=self.table.xview)
        horizontal.grid(row=1, column=0, sticky="ew")
        self.table.configure(yscrollcommand=vertical.set, xscrollcommand=horizontal.set)
        pane.add(results_frame, weight=2)
        self.folder.trace_add("write", self.folder_changed)
        root.protocol("WM_DELETE_WINDOW", self.close)
        self.after_id = root.after(80, self.poll)
        root.bind("<Destroy>", self.on_destroy, add="+")
        self.schedule_service(delay=0)

    def write(self, message):
        self.output.configure(state="normal")
        self.output.insert("end", str(message).replace("\x00", "\\0") + "\n")
        if int(self.output.index("end-1c").split(".")[0]) > 2500:
            self.output.delete("1.0", "501.0")
        self.output.see("end")
        self.output.configure(state="disabled")

    def clear_results(self):
        children = self.table.get_children()
        if children:
            self.table.delete(*children)
        self.last_report = None

    def folder_changed(self, *args):
        if not self.busy:
            self.clear_results()
            self.status.set("目录已变更，正在切换监听并后台核对索引。")
            self.schedule_service()

    def choose_folder(self):
        from tkinter import filedialog
        folder = filedialog.askdirectory(initialdir=self.folder.get() or "D:/")
        if folder:
            self.folder.set(folder)

    def new_database(self):
        self.database.set(str(fresh_database()))
        self.clear_results()
        self.status.set("已选择全新数据库，后台开始建库；原数据库保留。")
        self.schedule_service(delay=0)

    def schedule_service(self, delay=400):
        self.generation += 1
        if self.restart_id is not None:
            self.root.after_cancel(self.restart_id)
        if self.service:
            self.service.stop()
        self.live_status.set("正在切换目录/索引；搜索将等待后台同步。")
        self.restart_id = self.root.after(delay, self.restart_service)

    def restart_service(self):
        self.restart_id = None
        if self.closing:
            return
        if self.service and self.service.thread.is_alive():
            self.restart_id = self.root.after(40, self.restart_service)
            return
        folder = self.folder.get().strip()
        if not folder:
            self.live_status.set("请选择 Excel 文件夹。")
            return
        generation = self.generation
        try:
            self.service = LiveIndexService(folder, self.database.get(),
                lambda kind, value: self.events.put(("service", (generation, kind, value))))
            self.worker = self.service.thread
            self.service.start()
            if self.queued_request:
                action, options = self.queued_request
                self.queued_request = None
                self.service.request(action, **options)
        except Exception as exc:
            self.live_status.set(friendly_error(exc))
            self.status.set(friendly_error(exc))
            self.set_busy(False)

    def set_busy(self, value):
        self.busy = value
        for widget in self.inputs:
            widget.configure(state="disabled" if value else "normal")
        self.stop_button.configure(state="normal" if value else "disabled")

    def start(self, action):
        if self.busy or self.closing:
            return
        self.clear_results()
        try:
            folder = self.folder.get().strip()
            if not folder:
                raise ValueError("请选择 Excel 文件夹。")
            database = database_path(self.database.get())
            term = self.term.get()
            repeats = 1
            if action == "search":
                if not term.strip():
                    raise ValueError("请输入非空搜索关键词。")
                try:
                    repeats = int(self.repeat.get())
                except ValueError:
                    raise ValueError("重复次数必须为 1～1000 的整数。") from None
                if not 1 <= repeats <= 1000:
                    raise ValueError("重复次数必须为 1～1000 的整数。")
        except ValueError as exc:
            self.status.set(str(exc))
            self.write(str(exc))
            return
        compare = self.compare.get()
        self.set_busy(True)
        self.status.set("正在建立/更新索引…" if action == "build" else "等待后台同步完成后搜索…")
        self.write("\n" + ("建立/更新索引" if action == "build" else f"搜索：{term}"))

        options = dict(term=term, repeat=repeats, compare=compare)
        if self.service and not self.service.shutdown.is_set() and self.restart_id is None:
            self.service.request(action, **options)
        else:
            self.queued_request = (action, options)
            if self.restart_id is not None:
                self.root.after_cancel(self.restart_id)
            self.restart_id = self.root.after(0, self.restart_service)

    def stop(self):
        if self.service:
            self.service.cancel()
            self.status.set("正在停止；当前文件回滚，已经完成的文件索引保留。")
            self.stop_button.configure(state="disabled")

    def poll(self):
        self.after_id = None
        # Bound each poll so a long log queue cannot starve the UI.
        for _ in range(150):
            try:
                kind, value = self.events.get_nowait()
            except queue.Empty:
                break
            if kind == "service":
                generation, kind, value = value
                if generation != self.generation:
                    continue
            if kind == "log":
                self.write(value)
                self.progress.set(str(value)[:300])
            elif kind == "results":
                report, results = value
                self.last_report = report
                for result in results:
                    display = [str(cell).replace("\x00", "\\0")[:1200] for cell in result]
                    self.table.insert("", "end", values=display)
                self.write(f"表格展示前 {len(results)} 条；匹配计数与校验包含全部结果。")
            elif kind == "done":
                self.set_busy(False)
                self.status.set(value)
            elif kind == "live":
                self.live_status.set(value)
        if self.closing and (self.service is None or not self.service.thread.is_alive()):
            self.set_busy(False)
            self.root.destroy()
            return
        self.after_id = self.root.after(80, self.poll)

    def on_destroy(self, event):
        if event.widget is self.root:
            if self.service:
                self.service.stop()
            if self.restart_id is not None:
                self.root.after_cancel(self.restart_id)
                self.restart_id = None
            if self.after_id is not None:
                self.root.after_cancel(self.after_id)
                self.after_id = None

    def close(self):
        self.closing = True
        if self.restart_id is not None:
            self.root.after_cancel(self.restart_id)
            self.restart_id = None
        if self.service:
            self.service.stop()
        self.status.set("正在停止监听和后台任务并关闭…")
        if self.service is None or not self.service.thread.is_alive():
            self.set_busy(False)
            self.root.destroy()

def gui(folder="D:/", database=None):
    import tkinter as tk
    root = tk.Tk()
    window = SpeedTestWindow(root, folder, database)
    root.mainloop()
    return window

def main():
    parser = argparse.ArgumentParser(description="ExcelSearch 单文件版：自动增量建索引与查询测速")
    parser.add_argument("--folder", default="D:/", help="Excel 文件夹，默认 D:/")
    parser.add_argument("--db", type=Path, help="branch 内的数据库路径；默认复用固定数据库")
    parser.add_argument("--build", action="store_true", help="增量更新并重试失败文件")
    parser.add_argument("--keyword", help="搜索关键词；自动建立或更新索引")
    parser.add_argument("--repeat", type=int, default=1)
    parser.add_argument("--compare", action="store_true", help="扫描相同成功文件并核对全部结果")
    parser.add_argument("--gui", action="store_true")
    parser.add_argument("--fresh", action="store_true", help="分配新数据库，保留旧库，测试完整建库")
    args = parser.parse_args()
    try:
        database = database_path(args.db or (fresh_database() if args.fresh else DEFAULT_DATABASE))
    except ValueError as exc:
        parser.error(str(exc))
    if args.gui or not (args.build or args.keyword is not None):
        gui(args.folder, database)
        return 0
    try:
        if args.build:
            build_index(args.folder, database)
        if args.keyword is not None:
            report, _ = benchmark(args.folder, database, args.keyword, args.repeat, args.compare)
            print(json.dumps(report, ensure_ascii=False, indent=2))
            if report["comparison"] in ("mismatch", "source_changed_or_unreadable"):
                return 2
        return 0
    except KeyboardInterrupt:
        print("操作已停止，已提交的索引保留。", file=sys.stderr)
        return 130
    except Exception as exc:
        print("失败：" + friendly_error(exc), file=sys.stderr)
        return 1

if __name__ == "__main__":
    raise SystemExit(main())
