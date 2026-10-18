"""Verify search results against explicit expected values, independent of Git."""

import os
import sys
import tempfile
import threading
import unittest
from datetime import datetime, time, timedelta
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import ExcelSearch as current
import xlwt
from openpyxl import Workbook
from openpyxl.packaging.relationship import Relationship
from openpyxl.workbook.external_link.external import ExternalBook, ExternalLink, ExternalSheetNames


class SearchResultTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.path = self.root / "values.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "报价"
        sheet.append([None, "description", "RMB", "USD", "数量", "金额", "备注"])
        sheet.append([None, "Alpha 中文", 12.5, 2, 4, 50, "Alpha second match"])
        sheet.append([None, "Beta", 10, 1.5, 3, 30, "_x0041_"])
        sheet.append([True, False, "#DIV/0!", datetime(2024, 1, 1), time(12, 30), timedelta(days=2)])
        sheet["J40"] = "Alpha offset"
        sheet["D5"] = "='[1]External'!A1"
        sheet["D6"] = "=1+1"
        sheet["D7"] = "=1/0"
        sheet.merge_cells("B10:C10")
        sheet["B10"] = "Alpha merged"
        sheet.row_dimensions[2].hidden = True
        workbook.create_sheet("empty")
        hidden = workbook.create_sheet("hidden")
        hidden.sheet_state = "veryHidden"
        hidden["C35"] = "Alpha hidden"
        link = ExternalLink(externalBook=ExternalBook(
            sheetNames=ExternalSheetNames(sheetName=["External"]),
        ))
        link.file_link = Relationship(
            Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLinkPath",
            Target="file:///nonexistent/external.xlsx", TargetMode="External",
        )
        workbook._external_links.append(link)
        workbook.save(self.path)
        workbook.close()
        # openpyxl cannot calculate formulas; inject Excel-style saved results.
        with ZipFile(self.path) as archive:
            parts = [(item, archive.read(item.filename)) for item in archive.infolist()]
        with ZipFile(self.path, "w") as archive:
            for item, data in parts:
                if item.filename == "xl/worksheets/sheet1.xml":
                    text = data.decode()
                    text = text.replace("<f>'[1]External'!A1</f><v></v>", "<f>'[1]External'!A1</f><v>123.45</v>")
                    text = text.replace("<f>1+1</f><v></v>", "<f>1+1</f><v>2</v>")
                    text = text.replace('<c r="D7"><f>1/0</f><v></v></c>', '<c r="D7" t="e"><f>1/0</f><v>#DIV/0!</v></c>')
                    data = text.encode()
                archive.writestr(item, data)

    def search(self, module, term, keywords=None, path=None):
        results, error = module.search_file(
            str(path or self.path), term, keywords or module.DEFAULT_KEYWORDS, threading.Event(),
        )
        self.assertEqual(error, "")
        return results

    def test_xlsx_and_xlsm_known_values_and_addresses(self):
        expected = {
            "alpha": {("报价", "B2"), ("报价", "G2"), ("报价", "J40"),
                      ("报价", "B10"), ("hidden", "C35")},
            "中文": {("报价", "B2")}, "12.5": {("报价", "C2")},
            "10": {("报价", "C3")}, "true": {("报价", "A4")},
            "false": {("报价", "B4")},
            "#div/0!": {("报价", "C4"), ("报价", "D7")},
            "2024": {("报价", "D4")}, "12:30": {("报价", "E4")},
            "2 days": {("报价", "F4")}, "_x0041_": {("报价", "G3")},
            "123.45": {("报价", "D5")}, "no-match": set(),
        }
        for suffix in (".xlsx", ".xlsm"):
            path = self.path.with_suffix(suffix)
            if path != self.path:
                path.write_bytes(self.path.read_bytes())
            for term, cells in expected.items():
                with self.subTest(suffix=suffix, term=term):
                    matches = self.search(current, term, path=path)
                    self.assertEqual(len(matches), len(cells))
                    self.assertEqual({(r.sheet_name, r.cell_address) for r in matches}, cells)
                    self.assertTrue(all(term in r.cell_value.lower() for r in matches))

    def test_known_addresses_multiple_hits_prices_and_external_formula_cache(self):
        matches = self.search(current, "alpha")
        self.assertEqual({(r.sheet_name, r.cell_address) for r in matches}, {
            ("报价", "B2"), ("报价", "G2"), ("报价", "J40"),
            ("报价", "B10"), ("hidden", "C35"),
        })
        row = next(r for r in matches if r.cell_address == "B2")
        self.assertEqual((row.description, row.rmb_price, row.usd_price), ("Alpha 中文", "12.5/件", "2"))
        external = self.search(current, "123.45")
        self.assertEqual([(r.cell_address, r.cell_value) for r in external], [("D5", "123.45")])
        errors = self.search(current, "#div/0!")
        self.assertEqual({r.cell_address for r in errors}, {"C4", "D7"})

    def test_custom_headers_copy_expected_same_row_values(self):
        keywords = {"description": ["备注"], "rmb": ["金额"], "usd": []}
        matches = self.search(current, "alpha", keywords)
        self.assertEqual(len(matches), 5)
        for result in matches:
            expected = ("Alpha second match", "50", "") if result.cell_address in ("B2", "G2") else ("", "", "")
            self.assertEqual((result.description, result.rmb_price, result.usd_price), expected)

    def test_xls_known_text_numbers_and_booleans(self):
        path = self.root / "legacy.xls"
        book = xlwt.Workbook()
        sheet = book.add_sheet("报价")
        for r, row in enumerate([["description", "RMB", "USD", "数量", "金额"],
                                 ["Alpha 中文", 12.5, 2, 4, 50], [True, False, -1, 0, ""]]):
            for c, value in enumerate(row):
                sheet.write(r, c, value)
        book.add_sheet("empty")
        book.save(str(path))
        expected = {"alpha": {"A2"}, "中文": {"A2"}, "12.5": {"B2"},
                    "true": {"A3"}, "false": {"B3"}, "no-match": set()}
        for term, cells in expected.items():
            with self.subTest(term=term):
                matches = self.search(current, term, path=path)
                self.assertEqual(len(matches), len(cells))
                self.assertEqual({r.cell_address for r in matches}, cells)
                self.assertTrue(all(r.sheet_name == "报价" for r in matches))

    def test_timed_and_final_cancellation(self):
        event = threading.Event()
        clock = iter((0, 0, 0.6))
        consumed = []

        def rows():
            for index in range(1000):
                if index == 128:
                    event.set()
                consumed.append(index)
                yield ("Alpha",)

        with patch.object(current, "monotonic", side_effect=lambda: next(clock)):
            self.assertIsNone(current.search_rows("x", "s", rows(), "alpha", current.DEFAULT_KEYWORDS, event))
        self.assertEqual(len(consumed), 129)

        event.clear()

        def final_rows():
            yield ("Alpha",)
            event.set()

        self.assertIsNone(current.search_rows("x", "s", final_rows(), "alpha", current.DEFAULT_KEYWORDS, event))


if __name__ == "__main__":
    unittest.main()
