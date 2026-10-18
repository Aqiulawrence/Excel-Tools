"""Independent expected-result checks, including malformed worksheet dimensions."""

import os
import sys
import tempfile
import threading
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook

import ExcelSearch as search
import xlwt


class CompletenessTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Use QApplication so these checks can run together with the window tests.
        from PyQt6.QtWidgets import QApplication
        cls.app = QApplication.instance() or QApplication([])
        cls.app.setQuitOnLastWindowClosed(False)

    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.path = self.root / "prices.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "prices"
        sheet.append(["description", "RMB", "USD", "quantity", "amount"])
        sheet.append(["Alpha", 12.5, 2, 4, 50])
        sheet["J40"] = "Alpha far column"
        sheet["D70"] = "Alpha far row"
        workbook.save(self.path)
        workbook.close()

    def change_dimensions(self, reference):
        with ZipFile(self.path) as archive:
            parts = [(item, archive.read(item.filename)) for item in archive.infolist()]
        with ZipFile(self.path, "w") as archive:
            for item, data in parts:
                if item.filename == "xl/worksheets/sheet1.xml":
                    root = ET.fromstring(data)
                    dimension = root.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}dimension")
                    dimension.set("ref", reference)
                    data = ET.tostring(root)
                archive.writestr(item, data)

    def assert_expected_cells(self):
        results, error = search.search_file(str(self.path), "alpha", search.DEFAULT_KEYWORDS, threading.Event())
        self.assertEqual(error, "")
        self.assertEqual({result.cell_address for result in results}, {"A2", "J40", "D70"})

    def test_wrong_row_dimension_must_not_hide_cells(self):
        self.change_dimensions("A1:J2")
        self.assert_expected_cells()

    def test_wrong_column_dimension_must_not_hide_cells(self):
        self.change_dimensions("A1:B70")
        self.assert_expected_cells()

    def test_a1_only_dimension_must_not_hide_all_data(self):
        self.change_dimensions("A1")
        self.assert_expected_cells()

    def test_no_dimension_must_still_find_all_cells(self):
        with ZipFile(self.path) as archive:
            parts = [(item, archive.read(item.filename)) for item in archive.infolist()]
        with ZipFile(self.path, "w") as archive:
            for item, data in parts:
                if item.filename == "xl/worksheets/sheet1.xml":
                    root = ET.fromstring(data)
                    root.remove(root.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}dimension"))
                    data = ET.tostring(root)
                archive.writestr(item, data)
        self.assert_expected_cells()

    def test_overstated_dimension_does_not_pad_empty_rows(self):
        self.change_dimensions("A1:XFD1048576")
        self.assert_expected_cells()

    def test_dollar_prefix_is_a_valid_workbook_name(self):
        renamed = self.root / "$prices.xlsx"
        self.path.rename(renamed)
        worker = search.ExcelSearchWorker()
        worker.setup("alpha", [str(self.root)], 2, search.DEFAULT_KEYWORDS)
        files = list(worker.collect_files())
        self.assertEqual(files, [str(renamed)])

    def test_only_excel_lock_prefix_is_skipped(self):
        (self.root / "~$prices.xlsx").write_bytes(b"Excel lock file")
        worker = search.ExcelSearchWorker()
        worker.setup("alpha", [str(self.root)], 2, search.DEFAULT_KEYWORDS)
        self.assertEqual(list(worker.collect_files()), [str(self.path)])

    def test_inaccessible_directory_must_not_be_silently_ignored(self):
        def denied_walk(path, onerror=None):
            if onerror is not None:
                onerror(PermissionError(13, "Access denied", str(self.root / "restricted")))
            yield str(self.root), [], [self.path.name]

        worker = search.ExcelSearchWorker()
        worker.setup("alpha", [str(self.root)], 1, search.DEFAULT_KEYWORDS)
        with patch.object(search.os, "walk", side_effect=denied_walk):
            self.assertEqual(list(worker.collect_files()), [str(self.path)])
        self.assertTrue(getattr(worker, "scan_errors", []), "directory access failure was hidden")

    def test_missing_search_directory_is_reported(self):
        worker = search.ExcelSearchWorker()
        worker.setup("alpha", [str(self.root / "missing")], 1, search.DEFAULT_KEYWORDS)
        worker.run()
        self.assertEqual(worker.counts, (0, 0, 0))
        self.assertEqual(len(worker.scan_errors), 1)

    def test_formula_without_saved_result_is_not_calculated(self):
        workbook = Workbook()
        workbook.active["A1"] = '=\"Alpha\"'
        workbook.save(self.path)
        workbook.close()
        # Document the reader's limitation, not a promise to calculate formulas.
        results, error = search.search_file(str(self.path), "alpha", search.DEFAULT_KEYWORDS, threading.Event())
        self.assertEqual(error, "")
        self.assertEqual(results, [])

    def test_out_of_range_numeric_text_cannot_discard_workbook_matches(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["description", "RMB", "USD", "quantity", "amount"])
        sheet.append(["Alpha unusual", "1e1000000", 2, 2, 4])
        sheet.append(["Alpha normal", 12, 2, 2, 24])
        workbook.save(self.path)
        workbook.close()
        results, error = search.search_file(str(self.path), "alpha", search.DEFAULT_KEYWORDS, threading.Event())
        self.assertEqual(error, "")
        self.assertEqual([r.cell_address for r in results], ["A2", "A3"])
        self.assertEqual(results[0].rmb_price, "1e1000000")
        self.assertEqual(results[1].rmb_price, "12/件")

    def test_malformed_keyword_settings_are_safe(self):
        for value in (None, [], 1, "invalid"):
            with self.subTest(value=value):
                self.assertEqual(search.clean_keywords(value), search.DEFAULT_KEYWORDS)
        self.assertEqual(search.clean_keywords({"description": "desc, 品名", "rmb": None, "usd": 3}),
                         {"description": ["desc", "品名"], "rmb": [], "usd": []})

    def test_xls_dates_times_durations_booleans_and_errors_are_searchable(self):
        for date_1904 in (False, True):
            book = xlwt.Workbook()
            book.dates_1904 = date_1904
            sheet = book.add_sheet("types")
            sheet.write(0, 0, datetime(2024, 1, 2), xlwt.easyxf(num_format_str="YYYY-MM-DD"))
            sheet.write(1, 0, 0.5, xlwt.easyxf(num_format_str="hh:mm"))
            sheet.write(2, 0, 2.5, xlwt.easyxf(num_format_str="[h]:mm:ss"))
            sheet.write(3, 0, True)
            sheet.write(4, 0, False)
            sheet.row(5).set_cell_error(0, "#DIV/0!")
            path = self.root / f"typed-{date_1904}.xls"
            book.save(str(path))
            for term, address, value in (
                ("2024-01-02", "A1", "2024-01-02 00:00:00"),
                ("12:00:00", "A2", "12:00:00"),
                ("2 days", "A3", "2 days, 12:00:00"),
                ("true", "A4", "True"), ("false", "A5", "False"),
                ("#div/0!", "A6", "#DIV/0!"),
            ):
                with self.subTest(date_1904=date_1904, term=term):
                    results, error = search.search_file(str(path), term, search.DEFAULT_KEYWORDS, threading.Event())
                    self.assertEqual(error, "")
                    self.assertIn((address, value), [(r.cell_address, r.cell_value) for r in results])


if __name__ == "__main__":
    unittest.main()
