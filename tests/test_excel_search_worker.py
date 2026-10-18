"""Integration checks for the real spawned workers; no Excel UI or reporting."""

import os
import sys
import tempfile
import threading
import unittest
from concurrent.futures.process import BrokenProcessPool
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook
from PyQt6.QtCore import QCoreApplication

import ExcelSearch as search


class SearchWorkerTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QCoreApplication.instance() or QCoreApplication([])

    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.folder = Path(directory.name)
        for index in range(20):
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "prices"
            sheet.append(["description", "RMB", "USD", "quantity", "amount"])
            sheet.append([f"Alpha {index}", 12.5, 2, 4, 50])
            sheet.append([f"Beta {index}", 10, 1.5, 3, 30])
            workbook.create_sheet("offset")["D40"] = "Alpha offset"
            workbook.save(self.folder / f"prices-{index}.xlsx")
            workbook.close()
        (self.folder / "broken.xlsx").write_bytes(b"broken workbook")
        (self.folder / "~$temporary.xlsx").write_bytes(b"must be skipped")
        self.pool = search.SearchProcessPool()
        self.addCleanup(self.pool.shutdown)

    def run_search(self, keyword="alpha", workers=2, cancel_on_progress=False, keywords=None):
        worker = search.ExcelSearchWorker(self.pool)
        worker.setup(keyword, [str(self.folder)], workers, keywords or search.DEFAULT_KEYWORDS)
        records = []
        progress = []
        worker.batch_result_signal.connect(records.extend)
        worker.progress_signal.connect(lambda *counts: progress.append(counts))
        if cancel_on_progress:
            worker.progress_signal.connect(lambda *_: worker.stop())
        worker.run()
        self.assertEqual(worker.error, "")
        self.assertEqual(worker.counts[2], len(records))
        return worker, records, progress

    def expected(self, keyword, keywords=None):
        results = []
        for file in self.folder.glob("prices-*.xlsx"):
            matches, error = search.search_file(
                str(file), keyword, keywords or search.DEFAULT_KEYWORDS, threading.Event(),
            )
            self.assertEqual(error, "")
            results.extend(matches)
        return sorted(tuple(vars(item).values()) for item in results)

    def test_reuse_pool_with_different_keyword_and_keyword_settings(self):
        first, records, progress = self.run_search()
        self.assertEqual(first.counts, (20, 1, 40))
        self.assertEqual(progress[-1], (21, 21, 40))
        self.assertEqual(sorted(tuple(vars(item).values()) for item in records), self.expected("alpha"))
        executor = self.pool.executor
        keywords = {"description": ["quantity"], "rmb": ["amount"], "usd": []}
        second, records, _ = self.run_search("beta", keywords=keywords)
        self.assertIs(self.pool.executor, executor)
        self.assertEqual(second.counts, (20, 1, 20))
        self.assertEqual(sorted(tuple(vars(item).values()) for item in records), self.expected("beta", keywords))

    def test_cancel_drain_then_new_search(self):
        self.run_search()
        stopped, _, progress = self.run_search(cancel_on_progress=True)
        self.assertTrue(stopped.cancelled)
        self.assertLess(progress[-1][0], 21)
        finished, records, _ = self.run_search("beta")
        self.assertEqual(finished.counts, (20, 1, 20))
        self.assertEqual(sorted(tuple(vars(item).values()) for item in records), self.expected("beta"))

    def test_changing_worker_count_replaces_pool_and_shutdown_is_idempotent(self):
        self.run_search()
        executor = self.pool.executor
        worker, _, _ = self.run_search(workers=1)
        self.assertEqual(worker.counts, (20, 1, 40))
        self.assertIsNot(self.pool.executor, executor)
        self.pool.shutdown()
        self.pool.shutdown()
        self.assertIsNone(self.pool.executor)

    def test_empty_folder_and_standalone_worker_cleanup(self):
        empty = self.folder / "empty"
        empty.mkdir()
        worker = search.ExcelSearchWorker(self.pool)
        worker.setup("alpha", [str(empty)], 2, search.DEFAULT_KEYWORDS)
        worker.run()
        self.assertEqual(worker.counts, (0, 0, 0))
        self.assertIsNone(self.pool.executor)
        standalone = search.ExcelSearchWorker()
        standalone.setup("alpha", [str(self.folder)], 2, search.DEFAULT_KEYWORDS)
        standalone.run()
        self.assertEqual(standalone.counts, (20, 1, 40))
        self.assertIsNone(standalone.process_pool.executor)

    def test_idle_process_crash_does_not_break_next_search(self):
        self.run_search()
        with self.assertRaises(BrokenProcessPool):
            self.pool.executor.submit(os._exit, 1).result(timeout=10)
        worker, records, _ = self.run_search("beta")
        self.assertEqual(worker.counts, (20, 1, 20))
        self.assertEqual(sorted(tuple(vars(item).values()) for item in records), self.expected("beta"))

    def test_pool_reuse_reads_edits_additions_and_deletions(self):
        self.run_search()
        path = self.folder / "prices-0.xlsx"
        workbook = Workbook()
        workbook.active["C9"] = "Changed new content"
        workbook.save(path)
        workbook.close()
        (self.folder / "prices-1.xlsx").unlink()
        path = self.folder / "new.XLSX"
        workbook = Workbook()
        workbook.active["E11"] = "Changed added file"
        workbook.save(path)
        workbook.close()
        worker, records, _ = self.run_search("changed")
        self.assertEqual(worker.counts, (20, 1, 2))
        self.assertEqual({(Path(r.file_path).name, r.cell_address) for r in records},
                         {("prices-0.xlsx", "C9"), ("new.XLSX", "E11")})

    def test_parallelism_does_not_lose_or_duplicate_matches(self):
        for workers in (1, 4, 12, 16):
            with self.subTest(workers=workers):
                worker, records, progress = self.run_search(workers=workers)
                expected = {(f"prices-{i}.xlsx", sheet, cell) for i in range(20)
                            for sheet, cell in (("prices", "A2"), ("offset", "D40"))}
                actual = [(Path(r.file_path).name, r.sheet_name, r.cell_address) for r in records]
                self.assertEqual(len(actual), len(expected))
                self.assertEqual(set(actual), expected)
                self.assertEqual(worker.counts, (20, 1, 40))
                self.assertEqual(progress[-1], (21, 21, 40))




if __name__ == "__main__":
    unittest.main()
